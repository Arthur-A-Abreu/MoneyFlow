import io
import csv
from datetime import date
from decimal import Decimal
from django.db.models import Sum
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

from apps.finance.models import Transaction


def get_filtered_transactions(user, date_from=None, date_to=None, category_id=None, transaction_type=None):
    """Retorna movimentações filtradas com segurança."""
    qs = Transaction.objects.for_user(user).select_related('category')
    if date_from:
        qs = qs.filter(date__gte=date_from)
    if date_to:
        qs = qs.filter(date__lte=date_to)
    if category_id:
        qs = qs.filter(category_id=category_id)
    if transaction_type:
        qs = qs.filter(transaction_type=transaction_type)
    return qs.order_by('date', 'created_at')


def export_transactions_csv(transactions):
    """Gera conteúdo CSV formatado das movimentações."""
    output = io.StringIO()
    writer = csv.writer(output, delimiter=';')
    
    # Cabeçalho
    writer.writerow(['Data', 'Tipo', 'Categoria', 'Descrição', 'Valor (R$)', 'Status', 'Observações'])
    
    for t in transactions:
        tipo = 'Receita' if t.transaction_type == 'INCOME' else 'Despesa'
        valor = f"{t.amount:.2f}".replace('.', ',')
        status = 'Concluída' if t.status == 'COMPLETED' else 'Pendente'
        writer.writerow([
            t.date.strftime('%d/%m/%Y'),
            tipo,
            t.category.name,
            t.description,
            valor,
            status,
            t.notes or ''
        ])
        
    return output.getvalue().encode('utf-8-sig')


def export_transactions_excel(transactions, user, date_from=None, date_to=None):
    """Gera uma planilha Excel estilizada e profissional (.xlsx)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Extrato Financeiro"

    # Paleta de Estilo
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E293B")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="64748B")
    
    income_font = Font(name="Calibri", size=11, color="059669", bold=True)
    expense_font = Font(name="Calibri", size=11, color="DC2626", bold=True)
    regular_font = Font(name="Calibri", size=11, color="1E293B")
    total_font = Font(name="Calibri", size=11, bold=True, color="0F172A")
    total_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color="E2E8F0"),
        right=Side(style='thin', color="E2E8F0"),
        top=Side(style='thin', color="E2E8F0"),
        bottom=Side(style='thin', color="E2E8F0")
    )

    # Título do Relatório
    ws['A1'] = "MoneyFlow — Relatório Financeiro"
    ws['A1'].font = title_font
    ws.merge_cells('A1:G1')

    period_str = f"Período: {date_from.strftime('%d/%m/%Y') if date_from else 'Início'} até {date_to.strftime('%d/%m/%Y') if date_to else 'Atual'}"
    ws['A2'] = f"Usuário: {user.get_full_name() or user.username} | {period_str} | Gerado em: {date.today().strftime('%d/%m/%Y')}"
    ws['A2'].font = subtitle_font
    ws.merge_cells('A2:G2')

    headers = ["Data", "Tipo", "Categoria", "Descrição", "Valor (R$)", "Status", "Observações"]
    for col_num, header_title in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_num, value=header_title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_num = 5
    total_income = Decimal('0')
    total_expense = Decimal('0')

    for t in transactions:
        ws.cell(row=row_num, column=1, value=t.date.strftime('%d/%m/%Y')).alignment = Alignment(horizontal="center")
        
        tipo_str = "Receita" if t.transaction_type == 'INCOME' else "Despesa"
        ws.cell(row=row_num, column=2, value=tipo_str).alignment = Alignment(horizontal="center")
        
        ws.cell(row=row_num, column=3, value=t.category.name)
        ws.cell(row=row_num, column=4, value=t.description)
        
        val_cell = ws.cell(row=row_num, column=5, value=float(t.amount))
        val_cell.number_format = 'R$ #,##0.00'
        if t.transaction_type == 'INCOME':
            val_cell.font = income_font
            total_income += t.amount
        else:
            val_cell.font = expense_font
            total_expense += t.amount
            
        status_str = "Concluída" if t.status == 'COMPLETED' else "Pendente"
        ws.cell(row=row_num, column=6, value=status_str).alignment = Alignment(horizontal="center")
        ws.cell(row=row_num, column=7, value=t.notes or "")

        for c in range(1, 8):
            ws.cell(row=row_num, column=c).border = thin_border
            if c not in (5,):
                ws.cell(row=row_num, column=c).font = regular_font
        row_num += 1

    # Linhas de Totais
    row_num += 1
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=4)
    cell_tot_inc = ws.cell(row=row_num, column=1, value="Total de Receitas:")
    cell_tot_inc.font = total_font
    cell_tot_inc.alignment = Alignment(horizontal="right")
    val_tot_inc = ws.cell(row=row_num, column=5, value=float(total_income))
    val_tot_inc.font = income_font
    val_tot_inc.number_format = 'R$ #,##0.00'
    ws.cell(row=row_num, column=1).fill = total_fill
    ws.cell(row=row_num, column=5).fill = total_fill

    row_num += 1
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=4)
    cell_tot_exp = ws.cell(row=row_num, column=1, value="Total de Despesas:")
    cell_tot_exp.font = total_font
    cell_tot_exp.alignment = Alignment(horizontal="right")
    val_tot_exp = ws.cell(row=row_num, column=5, value=float(total_expense))
    val_tot_exp.font = expense_font
    val_tot_exp.number_format = 'R$ #,##0.00'
    ws.cell(row=row_num, column=1).fill = total_fill
    ws.cell(row=row_num, column=5).fill = total_fill

    row_num += 1
    saldo = total_income - total_expense
    ws.merge_cells(start_row=row_num, start_column=1, end_row=row_num, end_column=4)
    cell_bal = ws.cell(row=row_num, column=1, value="Saldo Final do Período:")
    cell_bal.font = total_font
    cell_bal.alignment = Alignment(horizontal="right")
    val_bal = ws.cell(row=row_num, column=5, value=float(saldo))
    val_bal.font = income_font if saldo >= 0 else expense_font
    val_bal.number_format = 'R$ #,##0.00'
    ws.cell(row=row_num, column=1).fill = total_fill
    ws.cell(row=row_num, column=5).fill = total_fill

    # Ajuste de larguras de coluna
    col_widths = {'A': 12, 'B': 12, 'C': 20, 'D': 35, 'E': 16, 'F': 14, 'G': 30}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def export_transactions_pdf(transactions, user, date_from=None, date_to=None):
    """Gera um PDF elegante e formatado do extrato financeiro usando ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        textColor=colors.HexColor('#64748B'),
        spaceAfter=15,
    )
    cell_style = ParagraphStyle(
        'DocCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        textColor=colors.HexColor('#1E293B'),
    )
    header_cell_style = ParagraphStyle(
        'DocHeaderCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9,
        textColor=colors.white,
        alignment=1, # Center
    )

    elements = []

    # Título & Cabeçalho
    elements.append(Paragraph("MoneyFlow — Relatório Financeiro", title_style))
    period_str = f"Período: {date_from.strftime('%d/%m/%Y') if date_from else 'Início'} até {date_to.strftime('%d/%m/%Y') if date_to else 'Atual'}"
    elements.append(Paragraph(f"Usuário: <b>{user.get_full_name() or user.username}</b> &nbsp;|&nbsp; {period_str} &nbsp;|&nbsp; Gerado em: {date.today().strftime('%d/%m/%Y')}", subtitle_style))

    # Tabela de Movimentações
    data = [
        [
            Paragraph("Data", header_cell_style),
            Paragraph("Tipo", header_cell_style),
            Paragraph("Categoria", header_cell_style),
            Paragraph("Descrição", header_cell_style),
            Paragraph("Valor (R$)", header_cell_style),
            Paragraph("Status", header_cell_style),
        ]
    ]

    total_income = Decimal('0')
    total_expense = Decimal('0')

    for t in transactions:
        tipo_str = "Receita" if t.transaction_type == 'INCOME' else "Despesa"
        tipo_color = "#059669" if t.transaction_type == 'INCOME' else "#DC2626"
        prefix = "+" if t.transaction_type == 'INCOME' else "-"
        
        if t.transaction_type == 'INCOME':
            total_income += t.amount
        else:
            total_expense += t.amount

        tipo_p = Paragraph(f"<font color='{tipo_color}'><b>{tipo_str}</b></font>", cell_style)
        val_p = Paragraph(f"<font color='{tipo_color}'><b>{prefix} R$ {t.amount:,.2f}</b></font>", cell_style)
        
        data.append([
            Paragraph(t.date.strftime('%d/%m/%Y'), cell_style),
            tipo_p,
            Paragraph(t.category.name, cell_style),
            Paragraph(t.description, cell_style),
            val_p,
            Paragraph("Concluída" if t.status == 'COMPLETED' else "Pendente", cell_style),
        ])

    table = Table(data, colWidths=[65, 60, 100, 180, 85, 50])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#4F46E5')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
    ]))
    elements.append(table)
    elements.append(Spacer(1, 15))

    # Tabela Resumo de Totais
    saldo = total_income - total_expense
    saldo_color = "#059669" if saldo >= 0 else "#DC2626"
    summary_data = [
        [Paragraph("<b>Resumo Financeiro</b>", cell_style), Paragraph("", cell_style)],
        [Paragraph("Total de Receitas:", cell_style), Paragraph(f"<font color='#059669'><b>+ R$ {total_income:,.2f}</b></font>", cell_style)],
        [Paragraph("Total de Despesas:", cell_style), Paragraph(f"<font color='#DC2626'><b>- R$ {total_expense:,.2f}</b></font>", cell_style)],
        [Paragraph("Saldo Final:", cell_style), Paragraph(f"<font color='{saldo_color}'><b>R$ {saldo:,.2f}</b></font>", cell_style)],
    ]
    summary_table = Table(summary_data, colWidths=[200, 150])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#E2E8F0')),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#F1F5F9')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(summary_table)

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
